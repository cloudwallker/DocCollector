from __future__ import annotations

import os

from .base import BaseExtractor, ExtractionResult, HitLocation, ContentUnit


def _iter_block_items(document):
    """Yield Paragraph and Table objects in document order (top level)."""
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    from docx.oxml.ns import qn

    body = document.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield Table(child, document)


class DocxExtractor(BaseExtractor):
    """Covers body paragraphs and tables (including nested tables).

    Headers, footers and text boxes are not indexed. Table hits report 表格/行/列 so the location is meaningful.
    """

    file_type = "docx"

    def extract(self) -> ExtractionResult:
        try:
            size = os.path.getsize(self.file_path)
        except OSError as exc:
            return ExtractionResult(
                path=self.file_path, file_type="docx", size=0,
                content="", status="error", error=f"无法获取文件大小: {exc}",
            )
        try:
            from docx import Document

            doc = Document(self.file_path)
            units = [ContentUnit(text, kind, value) for text, kind, value in self._walk(doc)]
            content = "\n".join(unit.text for unit in units)
        except Exception as exc:  # noqa: BLE001 - surface any parse failure
            return ExtractionResult(
                path=self.file_path, file_type="docx", size=size,
                content="", status="error", error=f"DOCX 解析失败: {exc}",
            )
        if not content.strip():
            return ExtractionResult(
                path=self.file_path, file_type="docx", size=size,
                content="", status="no_text", error="文档无可提取文本",
            )
        return ExtractionResult(
            path=self.file_path, file_type="docx", size=size, content=content, status="ok",
            units=units,
        )

    def search(self, keyword: str) -> list[HitLocation]:
        hits: list[HitLocation] = []
        if not keyword:
            return hits
        try:
            from docx import Document

            doc = Document(self.file_path)
        except Exception:
            return hits

        kw_lower = keyword.lower()
        kw_len = len(keyword)
        for text, loc_type, loc_value in self._walk(doc):
            text_lower = text.lower()
            start = 0
            while True:
                idx = text_lower.find(kw_lower, start)
                if idx == -1:
                    break
                ctx_start = max(0, idx - 50)
                ctx_end = min(len(text), idx + kw_len + 50)
                hits.append(
                    HitLocation(
                        location_type=loc_type,
                        location_value=loc_value,
                        context_before=text[ctx_start:idx].strip(),
                        matched_text=text[idx:idx + kw_len],
                        context_after=text[idx + kw_len:ctx_end].strip(),
                    )
                )
                start = idx + kw_len
        return hits

    # -- internal walking -------------------------------------------------

    def _walk(self, doc):
        """Yield ``(text, location_type, location_value)`` in document order."""
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        para_idx = 0
        tbl_idx = 0
        for block in _iter_block_items(doc):
            if isinstance(block, Paragraph):
                para_idx += 1
                if block.text:
                    yield block.text, "paragraph", para_idx
            elif isinstance(block, Table):
                tbl_idx += 1
                yield from self._walk_table(block, tbl_idx, "")

    def _walk_table(self, table, tbl_idx: int, nest_suffix: str):
        label_prefix = f"表格{tbl_idx}{nest_suffix}"
        for row_no, row in enumerate(table.rows, start=1):
            seen = set()
            for col_no, cell in enumerate(row.cells, start=1):
                tc_id = id(cell._tc)
                if tc_id in seen:  # merged cell already emitted
                    continue
                seen.add(tc_id)
                cell_text = cell.text
                if cell_text and cell_text.strip():
                    yield cell_text, "cell", f"{label_prefix} 行{row_no} 列{col_no}"
                for nested_no, nested in enumerate(cell.tables, start=1):
                    nested_suffix = f"{nest_suffix}›行{row_no}列{col_no}表{nested_no}"
                    yield from self._walk_table(nested, tbl_idx, nested_suffix)
